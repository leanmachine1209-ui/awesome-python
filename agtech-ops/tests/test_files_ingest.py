import io
import json

from agtech_ops.ingest import (
    SUPPORTED_EXTENSIONS,
    ingest_file,
    parse_excel,
    parse_json_records,
    parse_text_document,
)
from agtech_ops.models import Source


def test_supported_extensions_cover_core_formats():
    for ext in [".csv", ".tsv", ".xlsx", ".json", ".txt", ".pdf", ".docx"]:
        assert ext in SUPPORTED_EXTENSIONS


def test_json_records():
    payload = json.dumps(
        [
            {"farm": "Green Acres", "asset": "North Herd", "date": "2026-06-26",
             "metric": "milk_yield_l", "value": 1610, "notes": "ok"},
        ]
    )
    events, errors = parse_json_records(payload)
    assert errors == []
    assert len(events) == 1
    assert events[0].asset == "North Herd"
    assert events[0].value == 1610.0


def test_json_wrapped_in_records_key():
    payload = json.dumps({"records": [
        {"farm": "F", "asset": "A", "date": "2026-06-26", "value": 1}
    ]})
    events, errors = parse_json_records(payload)
    assert errors == []
    assert len(events) == 1


def test_excel_roundtrip():
    import pandas as pd

    df = pd.DataFrame(
        {
            "farm": ["Green Acres"],
            "asset": ["South Field"],
            "date": ["2026-06-27"],
            "metric": ["soil_moisture_pct"],
            "value": [15],
            "notes": ["dry"],
        }
    )
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    events, errors = parse_excel(buf.getvalue())
    assert errors == []
    assert len(events) == 1
    assert events[0].asset == "South Field"


def test_text_document_chunks_and_infers_asset():
    text = (
        "2026-06-26 North Herd cow looks lame, needs vet.\n\n"
        "Feed Store is running low, please reorder.\n"
    )
    events, errors = parse_text_document(
        text, farm="Green Acres", known_assets=["North Herd", "Feed Store"]
    )
    assert errors == []
    assert len(events) == 2
    assert events[0].asset == "North Herd"
    assert events[1].asset == "Feed Store"


def test_registry_dispatch_csv():
    csv = b"farm,asset,date,value\nGreen Acres,North Herd,2026-06-26,5\n"
    events, errors = ingest_file("data.csv", csv)
    assert errors == []
    assert len(events) == 1
    assert events[0].source is Source.csv_partner


def test_registry_dispatch_txt_whatsapp_autodetected():
    wa = (
        b"[2026-06-26, 7:32 AM] Alice: North Herd cow is lame\n"
        b"[2026-06-26, 7:40 AM] Bob: South Field looks dry\n"
    )
    events, errors = ingest_file(
        "chat.txt", wa, farm="Green Acres", known_assets=["North Herd", "South Field"]
    )
    assert errors == []
    assert len(events) == 2
    assert events[0].source is Source.whatsapp


def test_registry_text_requires_farm():
    events, errors = ingest_file("notes.txt", b"some notes here")
    assert events == []
    assert any("farm name is required" in e for e in errors)


def test_registry_unsupported_extension():
    events, errors = ingest_file("photo.heic", b"\x00\x01")
    assert events == []
    assert any("unsupported file type" in e for e in errors)


def test_docx_extraction():
    import docx

    d = docx.Document()
    d.add_paragraph("North Herd vaccination completed today.")
    d.add_paragraph("Feed Store needs a reorder soon.")
    buf = io.BytesIO()
    d.save(buf)
    events, errors = ingest_file(
        "report.docx", buf.getvalue(), farm="Green Acres",
        known_assets=["North Herd", "Feed Store"],
    )
    assert errors == []
    assert len(events) == 2
